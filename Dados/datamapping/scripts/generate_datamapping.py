"""
Script de referência para geração de Data Mapping Excel.
Este script contém todas as funções utilitárias necessárias.
Adapte a seção __main__ para o projeto específico.

Uso:
    python generate_datamapping.py

Requisitos:
    pip install openpyxl
"""

import os
import shutil
from copy import copy
from datetime import datetime

import openpyxl


# ============================================================================
# CONFIGURAÇÃO - Ajuste para o seu projeto
# ============================================================================

# Caminho do template (relativo ao diretório deste script)
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(SCRIPT_DIR, 'template.xlsx')

# Exemplo preenchido para referência visual (NÃO usar como template):
#   os.path.join(SCRIPT_DIR, 'exemplo.xlsx')

OUT_DIR = r'.\output'

# Larguras de coluna padrão das abas de serviço
SERVICE_COL_WIDTHS = {
    'A': 18.0, 'B': 21.14, 'C': 18.14, 'D': 90.43,
    'E': 72.29, 'F': 18.0, 'G': 41.57, 'H': 66.29,
    'I': 72.71, 'J': 71.0
}

# Alturas de linha padrão
ROW_HEIGHTS = {
    'ec_title': 60.75,      # {error codes} título
    'ec_header': 18.75,     # {error codes} headers
    'sv_title': 29.25,      # Serviço título da rota
    'sv_labels': 21.0,      # Serviço REQUEST/RESPONSE
}


# ============================================================================
# FUNÇÕES UTILITÁRIAS
# ============================================================================

def copy_cell_style(src, dst):
    """Copy style from src cell to dst cell."""
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format


def get_template_styles(template_path):
    """Extract template cell styles for replication."""
    wb = openpyxl.load_workbook(template_path)

    styles = {}

    # {error codes} styles
    ws = wb['{error codes}']
    styles['ec_title'] = ws['C4']       # row 4 - título (merge C4:G4)
    styles['ec_header'] = ws['C5']      # row 5 - headers de coluna
    styles['ec_cell'] = ws['C6']        # row 6 - células de dados

    # Serviço styles
    ws = wb['Serviço']
    styles['sv_title'] = ws['A2']       # row 2 - título da rota (merge A2:J2)
    styles['sv_req'] = ws['A3']         # row 3 - "REQUEST" (merge A3:E3)
    styles['sv_resp'] = ws['F3']        # row 3 - "RESPONSE" (merge F3:J3)
    styles['sv_header'] = ws['A4']      # row 4 - headers de coluna
    styles['sv_cell'] = ws['H5']        # row 5 - células de dados

    # {errors} styles
    ws = wb['{errors}']
    styles['er_title'] = ws['A2']
    styles['er_req'] = ws['A3']
    styles['er_resp'] = ws['F3']
    styles['er_header'] = ws['A4']
    styles['er_cell'] = ws['H5']

    wb.close()
    return styles


# ============================================================================
# PREENCHIMENTO DAS ABAS
# ============================================================================

def fill_capa(ws, service_name):
    """Fill the Capa sheet with service-specific data."""
    ws['B1'] = service_name
    ws['C4'] = service_name
    ws['C5'] = 'CongonhasHUB'
    ws['C6'] = '[X] CongonhasHUB_API'
    ws['C7'] = '1.0'

    # Histórico de Alterações
    ws['B17'] = 'Luigi Neto'
    ws['D17'] = datetime.now()
    ws['D17'].number_format = 'DD/MM/YYYY'
    ws['E17'] = '1.0'
    ws['F17'] = 'Criação do documento'

    # Aprovações
    ws['B25'] = 'Luigi Neto'
    ws['D25'] = datetime.now()
    ws['D25'].number_format = 'DD/MM/YYYY'
    ws['E25'] = '1.0'
    ws['F25'] = 'Criação do Template'


def write_error_codes_sheet(ws, styles, service_name, errors):
    """
    Fill {error codes} sheet.

    Args:
        ws: Worksheet object
        styles: Dict com estilos extraídos do template
        service_name: Nome do microserviço
        errors: Lista de tuplas (httpCode, detail, message)
    """
    # Title in C4 (already merged C4:G4 from template)
    ws['C4'] = f'MS {service_name} - Error Codes'

    # Row heights from template
    ws.row_dimensions[4].height = ROW_HEIGHTS['ec_title']
    ws.row_dimensions[5].height = ROW_HEIGHTS['ec_header']

    # Headers in row 5
    headers = ['Código de Erro do Backend', 'Mensagem', 'httpCode', 'detail', 'message']
    for i, h in enumerate(headers):
        cell = ws.cell(row=5, column=3 + i, value=h)
        copy_cell_style(styles['ec_header'], cell)

    # Data rows starting at row 6
    for idx, (code, detail, message) in enumerate(errors):
        r = 6 + idx
        for c in range(3, 8):
            cell = ws.cell(row=r, column=c)
            copy_cell_style(styles['ec_cell'], cell)

        ws.cell(row=r, column=5, value=code)
        ws.cell(row=r, column=6, value=detail)
        ws.cell(row=r, column=7, value=message)


def write_service_sheet(ws, styles, routes):
    """
    Fill a service route sheet.

    Args:
        ws: Worksheet object
        styles: Dict com estilos extraídos do template
        routes: Lista de dicts, cada um com:
            - 'title': str (ex: '[POST] /auth/register - REQUEST RegisterRequest / RESPONSE UserProfile')
            - 'request': lista de tuplas (tipo, localização, propriedade, regras, exemplo)
            - 'response': lista de tuplas (tipo, localização, propriedade, regras, obrigatório)
    """
    r = 2  # start row

    for route in routes:
        # Title row - merge A:J
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        cell = ws.cell(row=r, column=1, value=route['title'])
        copy_cell_style(styles['sv_title'], cell)
        ws.row_dimensions[r].height = ROW_HEIGHTS['sv_title']
        r += 1

        # REQUEST / RESPONSE labels
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        cell = ws.cell(row=r, column=1, value='REQUEST')
        copy_cell_style(styles['sv_req'], cell)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
        cell = ws.cell(row=r, column=6, value='RESPONSE')
        copy_cell_style(styles['sv_resp'], cell)
        ws.row_dimensions[r].height = ROW_HEIGHTS['sv_labels']
        r += 1

        # Column headers
        req_headers = ['Tipo de Container', 'Localização/Objeto', 'Propriedade', 'Regras', 'Exemplo']
        resp_headers = ['Tipo de Container', 'Localização/Objeto', 'Propriedade', 'Regras', 'Obrigatório?']
        for i, h in enumerate(req_headers):
            cell = ws.cell(row=r, column=1 + i, value=h)
            copy_cell_style(styles['sv_header'], cell)
        for i, h in enumerate(resp_headers):
            cell = ws.cell(row=r, column=6 + i, value=h)
            copy_cell_style(styles['sv_header'], cell)
        r += 1

        # Data rows
        max_rows = max(len(route.get('request', [])), len(route.get('response', [])))
        for i in range(max_rows):
            # Style all 10 cells
            for c in range(1, 11):
                copy_cell_style(styles['sv_cell'], ws.cell(row=r, column=c))

            if i < len(route.get('request', [])):
                for j, val in enumerate(route['request'][i]):
                    if val:
                        ws.cell(row=r, column=1 + j, value=val)

            if i < len(route.get('response', [])):
                for j, val in enumerate(route['response'][i]):
                    if val:
                        ws.cell(row=r, column=6 + j, value=val)
            r += 1

        r += 1  # blank row between routes


def write_errors_sheet(ws, styles):
    """Fill {errors} sheet with standard FastAPI error formats."""
    r = 2

    # ---- HTTPException ----
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    cell = ws.cell(row=r, column=1, value='HTTPException - RESPONSE (detail)')
    copy_cell_style(styles['er_title'], cell)
    ws.row_dimensions[r].height = ROW_HEIGHTS['sv_title']
    r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    cell = ws.cell(row=r, column=1, value='REQUEST')
    copy_cell_style(styles['er_req'], cell)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
    cell = ws.cell(row=r, column=6, value='RESPONSE')
    copy_cell_style(styles['er_resp'], cell)
    ws.row_dimensions[r].height = ROW_HEIGHTS['sv_labels']
    r += 1

    req_h = ['Tipo de Container', 'Localização/Objeto', 'Propriedade', 'Regras', 'Exemplo']
    resp_h = ['Tipo de Container', 'Localização/Objeto', 'Propriedade', 'Regras', 'Obrigatório?']
    for i, h in enumerate(req_h):
        cell = ws.cell(row=r, column=1 + i, value=h)
        copy_cell_style(styles['er_header'], cell)
    for i, h in enumerate(resp_h):
        cell = ws.cell(row=r, column=6 + i, value=h)
        copy_cell_style(styles['er_header'], cell)
    r += 1

    # detail response
    for c in range(1, 11):
        copy_cell_style(styles['er_cell'], ws.cell(row=r, column=c))
    ws.cell(row=r, column=6, value='JSON Body')
    ws.cell(row=r, column=7, value='root')
    ws.cell(row=r, column=8, value='detail')
    ws.cell(row=r, column=9, value='Mensagem de erro')
    ws.cell(row=r, column=10, value='S')
    r += 2

    # ---- ValidationError (422) ----
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    cell = ws.cell(row=r, column=1, value='ValidationError (422) - RESPONSE (detail[])')
    copy_cell_style(styles['er_title'], cell)
    ws.row_dimensions[r].height = ROW_HEIGHTS['sv_title']
    r += 1

    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    cell = ws.cell(row=r, column=1, value='REQUEST')
    copy_cell_style(styles['er_req'], cell)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=10)
    cell = ws.cell(row=r, column=6, value='RESPONSE')
    copy_cell_style(styles['er_resp'], cell)
    ws.row_dimensions[r].height = ROW_HEIGHTS['sv_labels']
    r += 1

    for i, h in enumerate(req_h):
        cell = ws.cell(row=r, column=1 + i, value=h)
        copy_cell_style(styles['er_header'], cell)
    for i, h in enumerate(resp_h):
        cell = ws.cell(row=r, column=6 + i, value=h)
        copy_cell_style(styles['er_header'], cell)
    r += 1

    for prop, regra in [('loc', 'Localização do erro'), ('msg', 'Mensagem de validação'), ('type', 'Tipo de erro')]:
        for c in range(1, 11):
            copy_cell_style(styles['er_cell'], ws.cell(row=r, column=c))
        ws.cell(row=r, column=6, value='JSON Body')
        ws.cell(row=r, column=7, value='root.detail[]')
        ws.cell(row=r, column=8, value=prop)
        ws.cell(row=r, column=9, value=regra)
        ws.cell(row=r, column=10, value='S')
        r += 1


# ============================================================================
# GERAÇÃO DO ARQUIVO
# ============================================================================

def generate_service(styles, service_name, error_codes, route_sheets):
    """
    Generate one data mapping Excel file for a microservice.

    Args:
        styles: Dict com estilos extraídos via get_template_styles()
        service_name: Nome do microserviço (ex: 'auth-service')
        error_codes: Lista de tuplas (httpCode, detail, message)
        route_sheets: Dict { nome_da_aba: [rotas] }
            Cada rota: {
                'title': str,
                'request': [(tipo, loc, prop, regras, exemplo), ...],
                'response': [(tipo, loc, prop, regras, obrig), ...]
            }
    """
    os.makedirs(OUT_DIR, exist_ok=True)

    # Copy template
    out_path = os.path.join(OUT_DIR, f'CongonhasHUB_API - DataMapping - {service_name}_1.0.xlsx')
    shutil.copy2(TEMPLATE, out_path)

    wb = openpyxl.load_workbook(out_path)

    # 1. Fill Capa
    fill_capa(wb['Capa'], service_name)

    # 2. Fill {error codes}
    write_error_codes_sheet(wb['{error codes}'], styles, service_name, error_codes)

    # 3. Rename base "Serviço" sheet and create additional ones if needed
    sheet_names = list(route_sheets.keys())
    base_ws = wb['Serviço']
    first_name = sheet_names[0]
    base_ws.title = first_name
    write_service_sheet(base_ws, styles, route_sheets[first_name])

    for extra_name in sheet_names[1:]:
        new_ws = wb.create_sheet(title=extra_name, index=wb.sheetnames.index('{errors}'))
        # Set column widths from template
        for col, w in SERVICE_COL_WIDTHS.items():
            new_ws.column_dimensions[col].width = w
        write_service_sheet(new_ws, styles, route_sheets[extra_name])

    # 4. Fill {errors}
    write_errors_sheet(wb['{errors}'], styles)

    wb.save(out_path)
    wb.close()
    print(f'  ✅ {service_name} → {out_path}')


# ============================================================================
# EXEMPLO DE USO
# ============================================================================

if __name__ == '__main__':
    print('🗃️  Gerando Data Mappings...\n')

    styles = get_template_styles(TEMPLATE)

    # --- Exemplo: auth-service ---
    generate_service(
        styles=styles,
        service_name='auth-service',
        error_codes=[
            (401, 'invalid_credentials', 'Credenciais inválidas'),
            (404, 'user_not_found', 'Usuário não encontrado'),
            (409, 'email_already_exists', 'Email já cadastrado'),
        ],
        route_sheets={
            'register': [{
                'title': '[POST] /auth/register - REQUEST RegisterRequest / RESPONSE UserProfile',
                'request': [
                    ('JSON Body', 'root', 'email', 'string - Email válido, obrigatório', 'user@example.com'),
                    ('JSON Body', 'root', 'password', 'string - Mínimo 8 caracteres', '********'),
                    ('JSON Body', 'root', 'name', 'string - Nome completo', 'João Silva'),
                ],
                'response': [
                    ('JSON Body', 'root', 'id', 'string - UUID v4', 'S'),
                    ('JSON Body', 'root', 'email', 'string - Email do usuário', 'S'),
                    ('JSON Body', 'root', 'name', 'string - Nome do usuário', 'S'),
                    ('JSON Body', 'root', 'created_at', 'string - ISO 8601', 'S'),
                ]
            }],
            'login': [{
                'title': '[POST] /auth/login - REQUEST LoginRequest / RESPONSE TokenResponse',
                'request': [
                    ('JSON Body', 'root', 'email', 'string - Email válido', 'user@example.com'),
                    ('JSON Body', 'root', 'password', 'string - Senha do usuário', '********'),
                ],
                'response': [
                    ('JSON Body', 'root', 'access_token', 'string - JWT token', 'S'),
                    ('JSON Body', 'root', 'token_type', 'string - Tipo do token', 'S'),
                ]
            }],
        }
    )

    print('\n✅ Todos os Data Mappings gerados com sucesso!')
